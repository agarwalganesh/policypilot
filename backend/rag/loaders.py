"""
PolicyPilot — Document Loaders
Handles PDF (PyMuPDF) and DOCX files.
Preserves page numbers for accurate citations.
"""
from __future__ import annotations

import os
from pathlib import Path
from typing import List, Optional

import pymupdf as fitz  # PyMuPDF (modern import)
from docx import Document as DocxDocument
from langchain_core.documents import Document

from backend.config import get_logger

logger = get_logger(__name__)


class PDFLoader:
    """
    Load a PDF using PyMuPDF.
    Returns one Document per page with accurate page metadata.
    """

    def __init__(self, file_path: str):
        self.file_path = Path(file_path)
        if not self.file_path.exists():
            raise FileNotFoundError(f"PDF not found: {file_path}")

    def load(self) -> List[Document]:
        docs: List[Document] = []
        try:
            pdf = fitz.open(str(self.file_path))
        except Exception as exc:
            logger.error("pdf_open_failed", path=str(self.file_path), error=str(exc))
            raise ValueError(f"Cannot open PDF: {self.file_path.name}") from exc

        if pdf.page_count == 0:
            logger.warning("pdf_empty", path=str(self.file_path))
            pdf.close()
            return docs

        for page_num in range(pdf.page_count):
            page = pdf[page_num]
            text = page.get_text("text")

            if not text or not text.strip():
                logger.debug("pdf_page_empty", page=page_num + 1, path=str(self.file_path))
                continue

            docs.append(Document(
                page_content=text.strip(),
                metadata={
                    "source_file": self.file_path.name,
                    "page": page_num + 1,
                    "total_pages": pdf.page_count,
                    "file_type": "pdf",
                },
            ))

        pdf.close()
        logger.info("pdf_loaded", pages=len(docs), path=str(self.file_path))
        return docs


class DOCXLoader:
    """
    Load a DOCX file.
    Approximates page numbers using paragraph counting.
    """
    PARAGRAPHS_PER_PAGE = 25  # conservative estimate

    def __init__(self, file_path: str):
        self.file_path = Path(file_path)
        if not self.file_path.exists():
            raise FileNotFoundError(f"DOCX not found: {file_path}")

    def load(self) -> List[Document]:
        docs: List[Document] = []
        try:
            doc = DocxDocument(str(self.file_path))
        except Exception as exc:
            logger.error("docx_open_failed", path=str(self.file_path), error=str(exc))
            raise ValueError(f"Cannot open DOCX: {self.file_path.name}") from exc

        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        if not paragraphs:
            logger.warning("docx_empty", path=str(self.file_path))
            return docs

        # Group paragraphs into page-sized chunks for page-level metadata
        for i in range(0, len(paragraphs), self.PARAGRAPHS_PER_PAGE):
            page_paragraphs = paragraphs[i: i + self.PARAGRAPHS_PER_PAGE]
            text = "\n\n".join(page_paragraphs)
            approx_page = (i // self.PARAGRAPHS_PER_PAGE) + 1
            estimated_total = max(1, (len(paragraphs) + self.PARAGRAPHS_PER_PAGE - 1) // self.PARAGRAPHS_PER_PAGE)

            docs.append(Document(
                page_content=text,
                metadata={
                    "source_file": self.file_path.name,
                    "page": approx_page,
                    "total_pages": estimated_total,
                    "file_type": "docx",
                    "page_note": "approximate",
                },
            ))

        logger.info("docx_loaded", sections=len(docs), path=str(self.file_path))
        return docs


def load_document(file_path: str) -> List[Document]:
    """
    Auto-detect file type and load accordingly.
    Returns list of page-level Documents.
    """
    path = Path(file_path)
    ext = path.suffix.lower()

    if ext == ".pdf":
        return PDFLoader(file_path).load()
    elif ext in (".docx", ".doc"):
        return DOCXLoader(file_path).load()
    else:
        raise ValueError(f"Unsupported file type: {ext}. Supported: pdf, docx")
